#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document


def paragraph_style_name(paragraph: Any) -> str:
    style = getattr(paragraph, "style", None)
    if style is None:
        return ""
    return getattr(style, "name", "") or ""


def heading_level(style_name: str) -> int | None:
    normalized = style_name.strip().lower()
    if normalized == "heading 1":
        return 1
    if normalized == "heading 2":
        return 2
    if normalized == "heading 3":
        return 3
    return None


def build_outline(document: Document) -> dict[str, Any]:
    cover_paragraphs: list[dict[str, str]] = []
    sections: list[dict[str, Any]] = []
    stack: list[dict[str, Any]] = []
    first_heading_found = False

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        style_name = paragraph_style_name(paragraph)
        level = heading_level(style_name)

        if level is None and not first_heading_found:
            cover_paragraphs.append({"style": style_name, "text": text})
            continue

        if level is not None:
            first_heading_found = True
            node = {
                "title": text,
                "style": style_name,
                "level": level,
                "guidance": [],
                "children": [],
            }

            while stack and stack[-1]["level"] >= level:
                stack.pop()

            if stack:
                stack[-1]["children"].append(node)
            else:
                sections.append(node)
            stack.append(node)
            continue

        if stack:
            stack[-1]["guidance"].append(text)

    tables: list[dict[str, Any]] = []
    for index, table in enumerate(document.tables, start=1):
        rows: list[list[str]] = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        header = rows[0] if rows else []
        sample_rows = rows[1:3] if len(rows) > 1 else []
        tables.append(
            {
                "index": index,
                "rows": len(rows),
                "cols": len(table.columns),
                "header": header,
                "sample_rows": sample_rows,
            }
        )

    return {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "cover_paragraphs": cover_paragraphs,
        "sections": sections,
        "tables": tables,
    }


def render_markdown(schema: dict[str, Any], source_path: Path) -> str:
    lines = [
        "# System Overview Template Guide",
        "",
        f"- Source: `{source_path.as_posix()}`",
        f"- Generated: `{schema['generated_at']}`",
        "",
        "## Cover Paragraphs",
        "",
    ]

    cover_paragraphs = schema.get("cover_paragraphs", [])
    if not cover_paragraphs:
        lines.append("- _None found._")
    else:
        for item in cover_paragraphs:
            lines.append(f"- `{item['style'] or 'default'}`: {item['text']}")

    lines.extend(["", "## Sections", ""])

    def emit(nodes: list[dict[str, Any]], depth: int = 0) -> None:
        for node in nodes:
            indent = "  " * depth
            lines.append(f"{indent}- L{node['level']} `{node['style']}` {node['title']}")
            for item in node.get("guidance", []):
                lines.append(f"{indent}  - guidance: {item}")
            emit(node.get("children", []), depth + 1)

    emit(schema.get("sections", []))

    lines.extend(["", "## Tables", ""])
    for table in schema.get("tables", []):
        lines.append(f"- Table {table['index']}: {table['rows']} rows x {table['cols']} cols")
        if table.get("header"):
            lines.append(f"  - header: {' | '.join(table['header'])}")
        for sample in table.get("sample_rows", []):
            lines.append(f"  - sample: {' | '.join(sample)}")

    lines.append("")
    return "\n".join(lines)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Extract a DOCX system overview template into AI-readable schema.")
    parser.add_argument("template", help="Path to the overview-design DOCX template.")
    parser.add_argument(
        "--output-dir",
        "-o",
        required=True,
        help="Directory where the JSON schema and markdown guide will be written.",
    )
    parser.add_argument(
        "--schema-name",
        default="template-schema.json",
        help="Output JSON filename. Default: template-schema.json",
    )
    parser.add_argument(
        "--guide-name",
        default="template-guide.md",
        help="Output markdown filename. Default: template-guide.md",
    )
    parser.add_argument("--overwrite", action="store_true", help="Overwrite existing outputs.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    template_path = Path(args.template)
    output_dir = Path(args.output_dir)

    if not template_path.exists():
        print(f"error: template not found: {template_path}", file=sys.stderr)
        return 1

    output_dir.mkdir(parents=True, exist_ok=True)
    schema_path = output_dir / args.schema_name
    guide_path = output_dir / args.guide_name

    if not args.overwrite:
        for path in (schema_path, guide_path):
            if path.exists():
                print(f"error: output already exists: {path}", file=sys.stderr)
                print("rerun with --overwrite to replace it", file=sys.stderr)
                return 1

    document = Document(str(template_path))
    schema = build_outline(document)
    schema["source"] = template_path.as_posix()

    schema_path.write_text(json.dumps(schema, ensure_ascii=False, indent=2), encoding="utf-8")
    guide_path.write_text(render_markdown(schema, template_path), encoding="utf-8")

    print(f"created {schema_path}")
    print(f"created {guide_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
