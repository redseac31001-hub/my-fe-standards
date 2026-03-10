#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.shared import Pt


SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_ROOT = SCRIPT_DIR.parent
DEFAULT_TEMPLATE_CONFIG_PATH = SKILL_ROOT / "assets" / "system-overview-template-config.json"
DEFAULT_TEMPLATE_PATH = SKILL_ROOT / "assets" / "templates" / "system-overview-template.docx"


def load_json(path: Path) -> dict[str, Any]:
    try:
        data = json.loads(path.read_text(encoding="utf-8-sig"))
    except FileNotFoundError as exc:
        raise ValueError(f"spec file not found: {path}") from exc
    except json.JSONDecodeError as exc:
        raise ValueError(f"spec file is not valid JSON: {exc}") from exc
    if not isinstance(data, dict):
        raise ValueError("spec root must be a JSON object")
    return data


def split_path(path_value: Any) -> list[str]:
    if isinstance(path_value, list):
        return [str(item) for item in path_value if str(item)]
    if isinstance(path_value, str):
        return [segment for segment in path_value.split(".") if segment]
    return []


def normalize_template_content_plan(config: dict[str, Any]) -> list[dict[str, Any]]:
    plan = config.get("content_plan")
    if not isinstance(plan, list):
        raise ValueError("template config missing content_plan list")

    normalized_plan: list[dict[str, Any]] = []
    for item in plan:
        if not isinstance(item, dict):
            continue
        normalized_item = {
            "title": item.get("title", ""),
            "style": item.get("style", "Normal"),
            "kind": item.get("kind"),
        }
        path_segments = split_path(item.get("path"))
        if path_segments:
            normalized_item["path"] = path_segments
        columns = item.get("columns")
        if isinstance(columns, list):
            normalized_columns: list[tuple[str, str]] = []
            for column in columns:
                if not isinstance(column, dict):
                    continue
                key = str(column.get("key", "")).strip()
                label = str(column.get("label", "")).strip()
                if key and label:
                    normalized_columns.append((key, label))
            if normalized_columns:
                normalized_item["columns"] = normalized_columns
        normalized_plan.append(normalized_item)

    return normalized_plan


def get_nested(data: dict[str, Any], path: list[str]) -> Any:
    current: Any = data
    for key in path:
        if not isinstance(current, dict):
            return None
        current = current.get(key)
    return current


def get_field_value(data: dict[str, Any], path_value: Any) -> Any:
    path = split_path(path_value)
    if not path:
        return None
    return get_nested(data, path)


def coerce_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "；".join(coerce_text(item) for item in value if coerce_text(item))
    if isinstance(value, (int, float, bool)):
        return str(value)
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def list_of_strings(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, str):
        text = value.strip()
        return [text] if text else []
    if not isinstance(value, list):
        return []
    return [coerce_text(item) for item in value if coerce_text(item)]


def list_of_dicts(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        return []
    return [item for item in value if isinstance(item, dict)]


def style_exists(document: DocumentObject, style_name: str) -> bool:
    try:
        document.styles[style_name]
        return True
    except KeyError:
        return False


def safe_style(document: DocumentObject, preferred: str, fallback: str) -> str:
    return preferred if style_exists(document, preferred) else fallback


def resolve_body_font_size(document: DocumentObject) -> Pt:
    for style_name in ("正文1", "正文内容", "Normal"):
        if not style_exists(document, style_name):
            continue
        size = document.styles[style_name].font.size
        if size is None:
            continue
        if size.pt <= 12.0:
            return size
    return Pt(12)


def apply_body_run_format(document: DocumentObject, run: Any) -> None:
    run.font.size = resolve_body_font_size(document)


def replace_paragraph_text(paragraph: Any, text: str) -> None:
    paragraph.text = text


def resolve_mapping_value(spec: dict[str, Any], mapping: dict[str, Any]) -> str:
    value = get_field_value(spec, mapping.get("source"))
    text = coerce_text(value)
    if mapping.get("format") == "wrap_full_width_parentheses" and text:
        return f"（{text}）"
    return text


def fill_cover(document: DocumentObject, spec: dict[str, Any], template_config: dict[str, Any] | None = None) -> None:
    cover = spec.get("cover", {}) if isinstance(spec.get("cover"), dict) else {}
    cover_config = {}
    if isinstance(template_config, dict) and isinstance(template_config.get("cover"), dict):
        cover_config = template_config.get("cover", {})
    non_empty_before_toc = []
    for paragraph in document.paragraphs:
        style_name = getattr(getattr(paragraph, "style", None), "name", "") or ""
        if style_name.lower().startswith("toc"):
            break
        if paragraph.text.strip():
            non_empty_before_toc.append(paragraph)

    paragraph_mappings = cover_config.get("paragraph_fields_before_toc", [])
    if isinstance(paragraph_mappings, list) and paragraph_mappings:
        for paragraph, mapping in zip(non_empty_before_toc, paragraph_mappings):
            if not isinstance(mapping, dict):
                continue
            value = resolve_mapping_value(spec, mapping)
            if value:
                replace_paragraph_text(paragraph, value)
    else:
        cover_values = [
            cover.get("document_title", ""),
            cover.get("document_code", ""),
            f"（{cover.get('template_version_label', '')}）" if cover.get("template_version_label") else "",
            cover.get("company_name", ""),
            cover.get("document_date", "")
        ]
        for paragraph, value in zip(non_empty_before_toc[:5], cover_values):
            if value:
                replace_paragraph_text(paragraph, value)

    metadata_table = cover_config.get("metadata_table")
    if isinstance(metadata_table, dict) and len(document.tables) > int(metadata_table.get("table_index", -1)):
        table = document.tables[int(metadata_table.get("table_index", 0))]
        for mapping in metadata_table.get("cells", []):
            if not isinstance(mapping, dict):
                continue
            row = int(mapping.get("row", -1))
            col = int(mapping.get("col", -1))
            if row < 0 or col < 0 or row >= len(table.rows) or col >= len(table.columns):
                continue
            table.cell(row, col).text = resolve_mapping_value(spec, mapping)
    elif len(document.tables) >= 1:
        table = document.tables[0]
        table.cell(0, 1).text = coerce_text(cover.get("document_title"))
        table.cell(1, 1).text = coerce_text(cover.get("initial_version"))
        table.cell(2, 1).text = coerce_text(cover.get("confidentiality_level"))
        table.cell(2, 3).text = coerce_text(cover.get("document_version"))
        table.cell(3, 1).text = coerce_text(cover.get("author"))
        table.cell(3, 3).text = coerce_text(cover.get("author_date"))
        table.cell(4, 1).text = coerce_text(cover.get("reviewer"))
        table.cell(4, 3).text = coerce_text(cover.get("review_date"))

    version_control_table = cover_config.get("version_control_table")
    if isinstance(version_control_table, dict) and len(document.tables) > int(version_control_table.get("table_index", -1)):
        version_rows = list_of_dicts(get_field_value(spec, version_control_table.get("rows_field")))
        table = document.tables[int(version_control_table.get("table_index", 1))]
        keep_header_rows = max(int(version_control_table.get("keep_header_rows", 1)), 0)
        while len(table.rows) > keep_header_rows:
            table._tbl.remove(table.rows[-1]._tr)
        columns = version_control_table.get("columns", [])
        for row_data in version_rows:
            cells = table.add_row().cells
            for index, column in enumerate(columns):
                if index >= len(cells) or not isinstance(column, dict):
                    continue
                cells[index].text = coerce_text(row_data.get(column.get("source")))
    elif len(document.tables) >= 2:
        version_rows = list_of_dicts(spec.get("version_control"))
        table = document.tables[1]
        while len(table.rows) > 1:
            table._tbl.remove(table.rows[-1]._tr)
        for row_data in version_rows:
            cells = table.add_row().cells
            cells[0].text = coerce_text(row_data.get("version"))
            cells[1].text = coerce_text(row_data.get("date"))
            cells[2].text = coerce_text(row_data.get("participants"))
            cells[3].text = coerce_text(row_data.get("change_summary"))


def remove_body_from_first_heading(document: DocumentObject) -> None:
    body = document._element.body
    first_heading_element = None
    for paragraph in document.paragraphs:
        style_name = getattr(getattr(paragraph, "style", None), "name", "") or ""
        if style_name.strip().lower() == "heading 1":
            first_heading_element = paragraph._element
            break

    if first_heading_element is None:
        return

    removing = False
    for child in list(body):
        if child == first_heading_element:
            removing = True
        if removing and child.tag != qn("w:sectPr"):
            body.remove(child)


def add_heading_or_label(document: DocumentObject, title: str, style_name: str) -> None:
    paragraph = document.add_paragraph(style=safe_style(document, style_name, "Normal"))
    run = paragraph.add_run(title)
    if style_name not in {"heading 1", "heading 2", "heading 3", "Heading 1", "Heading 2", "Heading 3"}:
        apply_body_run_format(document, run)


def add_placeholder(document: DocumentObject) -> None:
    paragraph = document.add_paragraph(style=safe_style(document, "正文内容", "Normal"))
    run = paragraph.add_run("待补充")
    apply_body_run_format(document, run)


def add_bullets(document: DocumentObject, items: list[str]) -> None:
    if not items:
        add_placeholder(document)
        return
    bullet_style = safe_style(document, "List Paragraph", safe_style(document, "正文内容", "Normal"))
    for item in items:
        paragraph = document.add_paragraph(style=bullet_style)
        run = paragraph.add_run(item)
        apply_body_run_format(document, run)


def add_paragraphs(document: DocumentObject, items: list[str]) -> None:
    if not items:
        add_placeholder(document)
        return
    body_style = safe_style(document, "正文1", safe_style(document, "正文内容", "Normal"))
    for item in items:
        paragraph = document.add_paragraph(style=body_style)
        run = paragraph.add_run(item)
        apply_body_run_format(document, run)


def add_table(document: DocumentObject, rows: list[dict[str, Any]], columns: list[tuple[str, str]]) -> None:
    if not rows:
        add_placeholder(document)
        return
    table = document.add_table(rows=1, cols=len(columns))
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    header_cells = table.rows[0].cells
    for index, (_, label) in enumerate(columns):
        header_cells[index].text = label
        if header_cells[index].paragraphs:
            header_cells[index].paragraphs[0].style = safe_style(document, "正文1", "Normal")
            if header_cells[index].paragraphs[0].runs:
                apply_body_run_format(document, header_cells[index].paragraphs[0].runs[0])
    for row_data in rows:
        cells = table.add_row().cells
        for index, (key, _) in enumerate(columns):
            cells[index].text = coerce_text(row_data.get(key))
            if cells[index].paragraphs:
                cells[index].paragraphs[0].style = safe_style(document, "正文1", "Normal")
                if cells[index].paragraphs[0].runs:
                    apply_body_run_format(document, cells[index].paragraphs[0].runs[0])


def render_content(document: DocumentObject, spec: dict[str, Any], content_plan: list[dict[str, Any]]) -> None:
    for item in content_plan:
        title = item["title"]
        style = item["style"]
        path = item.get("path")
        kind = item.get("kind")

        add_heading_or_label(document, title, style)
        if kind is None:
            continue

        value = get_nested(spec, path) if isinstance(path, list) else None

        if kind == "bullets":
            add_bullets(document, list_of_strings(value))
        elif kind == "paragraphs":
            add_paragraphs(document, list_of_strings(value))
        elif kind == "table":
            add_table(document, list_of_dicts(value), item.get("columns", []))
        else:
            add_placeholder(document)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Render a system overview design DOCX from a template and JSON spec.")
    parser.add_argument(
        "--template",
        default=str(DEFAULT_TEMPLATE_PATH),
        help="Path to the overview-design DOCX template. Defaults to the bundled official template.",
    )
    parser.add_argument("--spec", required=True, help="Path to the structured JSON input.")
    parser.add_argument("--output", required=True, help="Output DOCX path.")
    parser.add_argument(
        "--template-config",
        default=str(DEFAULT_TEMPLATE_CONFIG_PATH),
        help="Template config JSON path. Defaults to the bundled system-overview template config.",
    )
    parser.add_argument("--overwrite", action="store_true", help="Overwrite the output file if it exists.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    template_path = Path(args.template)
    spec_path = Path(args.spec)
    output_path = Path(args.output)
    template_config_path = Path(args.template_config)

    if not template_path.exists():
        print(f"error: template not found: {template_path}", file=sys.stderr)
        return 1

    if output_path.exists() and not args.overwrite:
        print(f"error: output file already exists: {output_path}", file=sys.stderr)
        print("rerun with --overwrite to replace it", file=sys.stderr)
        return 1

    try:
        spec = load_json(spec_path)
        template_config = load_json(template_config_path)
        content_plan = normalize_template_content_plan(template_config)
    except ValueError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1

    document = Document(str(template_path))
    fill_cover(document, spec, template_config)
    remove_body_from_first_heading(document)
    render_content(document, spec, content_plan)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    print(f"created {output_path}")
    print(f"template config: {template_config_path}")
    print("note: open the document in Word and refresh the table of contents before final delivery")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
